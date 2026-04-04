# -*- coding: utf-8 -*-
"""優先順位仕様書.xlsx を生成する。"""
import os
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "優先順位仕様書.xlsx")
HEAD_FILL = PatternFill("solid", fgColor="E2EFDA")


def write_table(ws, start_row, headers, rows, col_widths=None):
    for j, h in enumerate(headers, 1):
        c = ws.cell(start_row, j, h)
        c.font = Font(bold=True)
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, row in enumerate(rows, 1):
        for j, v in enumerate(row, 1):
            ws.cell(start_row + i, j, v).alignment = Alignment(wrap_text=True, vertical="top")
    if col_widths:
        for j, w in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
    return start_row + len(rows) + 2


def main():
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "表紙"
    ws0["A1"] = "配台計画システム — 優先順位・解決ルール 仕様書"
    ws0["A1"].font = Font(size=16, bold=True)
    ws0["A3"] = "対象"
    ws0["B3"] = "planning_core.py（段階2）"
    ws0["A4"] = "出力日"
    ws0["B4"] = date.today().isoformat()
    ws0.column_dimensions["A"].width = 28
    ws0.column_dimensions["B"].width = 72

    ws4 = wb.create_sheet("04_タスク並べ替え")
    ws4["A1"] = "task_queue.sort のキー順"
    ws4["A1"].font = Font(bold=True)
    ws4.merge_cells("A1:D1")
    rows4 = [
        ["1", "due_source_rank", "納期ソース優先（小さいほど先）", "0:指定納期_上書き→1:AI完了/出荷→2:AI開始日→3:回答納期→4:指定納期→9:なし"],
        ["2", "0 if has_done_deadline_override else 1", "特別指定起因の締切圧縮を先行", "備考の締切意図を優先"],
        ["3", "0 if (in_progress and not has_special_remark) else 1", "加工途中かつ備考空を先行", "備考あり途中品は通常ソート"],
        ["4", "priority", "優先度（小さいほど先）", "セル優先で決定"],
        ["5", "0 if due_urgent else 1", "基準日以前期限を先に", "due_basis <= run_date"],
        ["6", "due_basis_date or date.max", "納期が早い順", "None は後ろ"],
        ["7", "start_date_req", "開始希望日が早い順", "原反投入など反映後"],
        ["8", "-resolve_need_required_op(...)", "必要人数が多い順", "tie-break"],
    ]
    write_table(ws4, 3, ["キー順", "式・条件", "意味", "備考"], rows4, [8, 40, 34, 44])

    ws6 = wb.create_sheet("06_AI備考解析_概要")
    ws6["A1"] = "特別指定_備考 → Gemini → 依頼NOキーJSON"
    ws6["A1"].font = Font(bold=True)
    ws6.merge_cells("A1:B1")
    rows6 = [
        ["解析対象", "依頼NO と 特別指定_備考 が揃う行のみ。依頼NOは 12345 / 12345.0 のゆれを正規化。"],
        ["キャッシュ", "output/ai_remarks_cache.json（TTL 6 時間）。TASK_SPECIAL_v2 + 基準年 + 備考blob でキー化。"],
        ["実行ログ", "log/execution_log.txt・log/ai_task_special_*・log/planning_conflict_highlight.tsv 等（常に最新1世代を上書き）。"],
        ["API キー", "GEMINI_API_KEY が無い場合は解析スキップ（空 dict）。"],
        ["反映", "列値があれば列優先、空欄のみAI補完。"],
    ]
    write_table(ws6, 3, ["項目", "内容"], rows6, [18, 86])

    wb.save(OUT)
    print(f"作成しました: {OUT}")


if __name__ == "__main__":
    main()

